"""Getting text out of a CV file.

Deliberately boring. ``pypdf`` and ``python-docx`` are pure Python and install
without a compiler on a locked-down machine, which matters more here than the
last few percent of layout fidelity — the model reads prose, not columns.
"""

from __future__ import annotations

import re
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".rtf"}

_MULTISPACE_RE = re.compile(r"[ \t]{2,}")
_BLANKS_RE = re.compile(r"\n{3,}")
_RTF_CONTROL_RE = re.compile(r"\\[a-z]+-?\d* ?|[{}]", re.IGNORECASE)


class UnreadableCV(RuntimeError):
    """The file could not be read, and the message says what to do instead."""


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _from_pdf(path)
    elif suffix == ".docx":
        text = _from_docx(path)
    elif suffix == ".rtf":
        text = _RTF_CONTROL_RE.sub(" ", path.read_text(encoding="utf-8", errors="replace"))
    elif suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".doc":
        raise UnreadableCV(
            "The old .doc format cannot be read directly. Open it in Word and "
            "save as .docx or PDF, then upload that."
        )
    else:
        raise UnreadableCV(
            f"{suffix or 'That file type'} is not supported. Upload a PDF, DOCX, "
            "TXT or Markdown file."
        )

    cleaned = tidy(text)
    if len(cleaned) < 200:
        raise UnreadableCV(
            "Almost no text came out of that file. If it is a scanned PDF there is "
            "nothing to read — export a text-based PDF from the original document, "
            "or upload the DOCX."
        )
    return cleaned


def _from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise UnreadableCV("PDF support is missing. Reinstall the dependencies.") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        raise UnreadableCV(f"That PDF could not be opened: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise UnreadableCV(
                "That PDF is password protected. Remove the password and upload it again."
            ) from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            # One unreadable page should not lose the other nine.
            pages.append("")
    return "\n\n".join(pages)


def _from_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise UnreadableCV("DOCX support is missing. Reinstall the dependencies.") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise UnreadableCV(f"That DOCX could not be opened: {exc}") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    # Plenty of CVs put the entire employment history in a table.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    # PDF extraction turns bullet glyphs into a scatter of symbols.
    text = re.sub(r"[\u2022\u25aa\u25cf\u25e6\uf0b7\u00b7]", "- ", text)
    text = _MULTISPACE_RE.sub(" ", text)
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return _BLANKS_RE.sub("\n\n", text).strip()
